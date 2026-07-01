"""
Viz Bridge - Generates charts using matplotlib and saves them
Creates visualizations (bar, line, pie, scatter) and HTML dashboards
"""

import os
import json
import base64
from typing import Dict, List, Optional, Tuple
from pathlib import Path
from io import BytesIO
from datetime import datetime

import matplotlib
matplotlib.use("Agg")  # Non-interactive backend for server use
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker

# Try to import python-pptx for slide generation
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

# Try to import plotly for interactive charts
try:
    import plotly.express as px
    import plotly.graph_objects as go
    import plotly.io as pio
    PLOTLY_AVAILABLE = True
except ImportError:
    PLOTLY_AVAILABLE = False

# Cores da marca 42c
COR_42C_PRIMARIA = "#1a1a2e"
COR_42C_SECUNDARIA = "#e94560"
COR_42C_ACAcento = "#0f3460"
COR_42C_TEXTO = "#16213e"
CORES_42C = ["#e94560", "#0f3460", "#1a1a2e", "#16213e", "#533483",
             "#e94560", "#ff6b6b", "#4ecdc4", "#45b7d1", "#96ceb4"]


class VizBridge:
    """
    Bridge between LangGraph and matplotlib for visualization.
    Generates charts, comparison graphs, and HTML dashboards.
    """

    def __init__(self, output_dir: str = "/opt/projetos/hermes-unified/output/viz"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

        # Color palette
        self.colors = [
            "#4C72B0", "#DD8452", "#55A868", "#C44E52",
            "#8172B3", "#937860", "#DA8BC3", "#8C8C8C",
            "#CCB974", "#64B5CD"
        ]

    def _validate_data(self, data: dict, chart_type: str) -> Tuple[bool, str]:
        """Validate input data structure for the chart type."""
        if not isinstance(data, dict):
            return False, "Data must be a dictionary"

        if chart_type in ["bar", "line", "pie"]:
            if "labels" not in data or "values" not in data:
                return False, f"'{chart_type}' requires 'labels' and 'values' keys"
            if not isinstance(data["labels"], list) or not isinstance(data["values"], list):
                return False, "labels and values must be lists"
            if len(data["labels"]) != len(data["values"]):
                return False, "labels and values must have same length"
            if len(data["labels"]) == 0:
                return False, "Data cannot be empty"

        elif chart_type == "scatter":
            if "x" not in data or "y" not in data:
                return False, "'scatter' requires 'x' and 'y' keys"
            if not isinstance(data["x"], list) or not isinstance(data["y"], list):
                return False, "x and y must be lists"
            if len(data["x"]) != len(data["y"]):
                return False, "x and y must have same length"

        return True, "Valid"

    def generate_chart(
        self,
        data: dict,
        chart_type: str,
        title: str,
        filename: Optional[str] = None,
        xlabel: str = "",
        ylabel: str = ""
    ) -> Tuple[str, bool, str]:
        """
        Generate a chart and save it as PNG.

        Args:
            data: {"labels": [...], "values": [...]} or {"x": [...], "y": [...]}
            chart_type: "bar", "line", "pie", "scatter"
            title: Chart title
            filename: Optional output filename (without path)
            xlabel: X-axis label (bar, line, scatter)
            ylabel: Y-axis label (bar, line, scatter)

        Returns:
            Tuple[filepath, success, message]
        """
        valid, msg = self._validate_data(data, chart_type)
        if not valid:
            return "", False, msg

        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            sanitized_title = "".join(c if c.isalnum() or c in " _-" else "_" for c in title)[:40]
            filename = f"{chart_type}_{sanitized_title}_{timestamp}.png"

        filepath = os.path.join(self.output_dir, filename)

        try:
            fig, ax = plt.subplots(figsize=(10, 6))

            if chart_type == "bar":
                ax.bar(data["labels"], data["values"], color=self.colors[:len(data["labels"])])
                ax.set_xlabel(xlabel or "Categories")
                ax.set_ylabel(ylabel or "Values")

            elif chart_type == "line":
                ax.plot(data["labels"], data["values"], marker="o", linewidth=2,
                        color=self.colors[0], markersize=6)
                ax.set_xlabel(xlabel or "X")
                ax.set_ylabel(ylabel or "Y")
                ax.grid(True, alpha=0.3)

            elif chart_type == "pie":
                wedges, texts, autotexts = ax.pie(
                    data["values"],
                    labels=data["labels"],
                    autopct="%1.1f%%",
                    colors=self.colors[:len(data["labels"])],
                    startangle=90
                )
                ax.axis("equal")

            elif chart_type == "scatter":
                scatter = ax.scatter(
                    data["x"], data["y"],
                    c=self.colors[0], alpha=0.7, s=80,
                    edgecolors="white", linewidth=0.5
                )
                ax.set_xlabel(xlabel or "X")
                ax.set_ylabel(ylabel or "Y")
                ax.grid(True, alpha=0.3)

            ax.set_title(title, fontsize=14, pad=15)
            fig.tight_layout()

            plt.savefig(filepath, dpi=150, bbox_inches="tight")
            plt.close(fig)

            return filepath, True, f"Chart saved: {filename}"

        except Exception as e:
            plt.close("all")
            return "", False, f"Chart generation error: {str(e)}"

    def generate_comparison_chart(
        self,
        datasets: List[dict],
        title: str,
        chart_type: str = "bar",
        filename: Optional[str] = None,
        xlabel: str = "",
        ylabel: str = ""
    ) -> Tuple[str, bool, str]:
        """
        Generate a comparison chart with multiple data series.

        Args:
            datasets: List of {"label": str, "values": [...]}
            title: Chart title
            chart_type: "bar" or "line"
            filename: Optional output filename
            xlabel, ylabel: Axis labels

        Returns:
            Tuple[filepath, success, message]
        """
        if not datasets:
            return "", False, "No datasets provided"
        if not all("label" in d and "values" in d for d in datasets):
            return "", False, "Each dataset needs 'label' and 'values' keys"

        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"comparison_{chart_type}_{timestamp}.png"

        filepath = os.path.join(self.output_dir, filename)

        try:
            fig, ax = plt.subplots(figsize=(12, 6))
            num_datasets = len(datasets)
            labels = datasets[0].get("labels", list(range(len(datasets[0]["values"]))))

            if chart_type == "bar":
                bar_width = 0.8 / num_datasets
                for i, ds in enumerate(datasets):
                    offset = (i - num_datasets / 2 + 0.5) * bar_width
                    x_pos = [j + offset for j in range(len(labels))]
                    bars = ax.bar(
                        x_pos, ds["values"],
                        width=bar_width * 0.9,
                        label=ds["label"],
                        color=self.colors[i % len(self.colors)]
                    )
                ax.set_xticks(range(len(labels)))
                ax.set_xticklabels(labels)

            elif chart_type == "line":
                for i, ds in enumerate(datasets):
                    ax.plot(
                        labels, ds["values"],
                        marker="o", linewidth=2,
                        label=ds["label"],
                        color=self.colors[i % len(self.colors)],
                        markersize=6
                    )
                ax.grid(True, alpha=0.3)

            ax.set_title(title, fontsize=14, pad=15)
            ax.set_xlabel(xlabel or "Categories")
            ax.set_ylabel(ylabel or "Values")
            ax.legend(loc="best")
            fig.tight_layout()

            plt.savefig(filepath, dpi=150, bbox_inches="tight")
            plt.close(fig)

            return filepath, True, f"Comparison chart saved: {filename}"

        except Exception as e:
            plt.close("all")
            return "", False, f"Comparison chart error: {str(e)}"

    def generate_html_dashboard(
        self,
        charts: List[str],
        title: str,
        filename: Optional[str] = None
    ) -> Tuple[str, bool, str]:
        """
        Create an HTML page with multiple embedded charts.

        Args:
            charts: List of PNG file paths to embed
            title: Dashboard title
            filename: Output filename

        Returns:
            Tuple[filepath, success, message]
        """
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"dashboard_{timestamp}.html"

        filepath = os.path.join(self.output_dir, filename)

        try:
            images_html = ""
            for i, chart_path in enumerate(charts):
                if os.path.exists(chart_path):
                    with open(chart_path, "rb") as f:
                        img_data = base64.b64encode(f.read()).decode()
                    chart_name = os.path.basename(chart_path)
                    images_html += f"""
                    <div class="chart-container">
                        <h2>Chart {i+1}: {chart_name}</h2>
                        <img src="data:image/png;base64,{img_data}"
                             alt="{chart_name}" style="max-width:100%;height:auto;">
                    </div>
                    <hr>
                    """
                else:
                    images_html += f"""
                    <div class="chart-container">
                        <p style="color:red;">Chart not found: {chart_path}</p>
                    </div>
                    """

            html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            max-width: 1200px; margin: 0 auto; padding: 20px;
            background: #f5f5f5; color: #333;
        }}
        h1 {{ color: #2c3e50; border-bottom: 2px solid #3498db; padding-bottom: 10px; }}
        .chart-container {{
            background: white; border-radius: 8px; padding: 20px;
            margin: 20px 0; box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }}
        .chart-container h2 {{ color: #7f8c8d; font-size: 16px; margin-top: 0; }}
        .footer {{ text-align: center; color: #95a5a6; font-size: 12px; margin-top: 30px; }}
        hr {{ border: none; border-top: 1px solid #ecf0f1; }}
        img {{ display: block; margin: 10px auto; }}
    </style>
</head>
<body>
    <h1>{title}</h1>
    <p class="meta">Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>
    {images_html}
    <div class="footer">
        <p>Generated by Hermes-Unified Viz Agent</p>
    </div>
</body>
</html>"""

            with open(filepath, "w") as f:
                f.write(html_content)

            return filepath, True, f"Dashboard saved: {filename}"

        except Exception as e:
            return "", False, f"Dashboard error: {str(e)}"

    def generate_pptx(
        self,
        charts: List[str],
        title: str,
        filename: Optional[str] = None
    ) -> Tuple[str, bool, str]:
        """
        Generate a PowerPoint presentation with embedded charts.

        Args:
            charts: List of PNG file paths to embed
            title: Presentation title
            filename: Output filename

        Returns:
            Tuple[filepath, success, message]
        """
        if not PPTX_AVAILABLE:
            return "", False, "python-pptx not installed. Install with: pip install python-pptx"

        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{title.replace(' ', '_')}_{timestamp}.pptx"

        filepath = os.path.join(self.output_dir, filename)

        try:
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)

            # Title slide
            title_slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
            from pptx.util import Inches, Pt
            txBox = title_slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(2))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(40)
            p.font.bold = True
            p.alignment = 2  # Center

            # Chart slides
            for chart_path in charts:
                if os.path.exists(chart_path):
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    chart_name = os.path.basename(chart_path)

                    # Title
                    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.6))
                    tf = txBox.text_frame
                    p = tf.paragraphs[0]
                    p.text = chart_name
                    p.font.size = Pt(18)

                    # Image
                    slide.shapes.add_picture(
                        chart_path,
                        Inches(0.5), Inches(1.2),
                        Inches(12.3), Inches(6)
                    )

            prs.save(filepath)
            return filepath, True, f"Presentation saved: {filename}"

        except Exception as e:
            return "", False, f"PPTX generation error: {str(e)}"

    def list_charts(self) -> List[dict]:
        """List all generated charts in the output directory."""
        charts = []
        for f in sorted(os.listdir(self.output_dir)):
            if f.endswith((".png", ".html", ".pptx")):
                fpath = os.path.join(self.output_dir, f)
                charts.append({
                    "filename": f,
                    "path": fpath,
                    "size": os.path.getsize(fpath),
                    "modified": datetime.fromtimestamp(os.path.getmtime(fpath)).isoformat()
                })
        return charts
    
    # =====================================================================
    # PLOTLY - Graficos interativos
    # =====================================================================
    
    def generate_plotly_chart(self, data: dict, chart_type: str = "bar", 
                               title: str = "Grafico") -> Tuple[str, bool, str]:
        """Generate interactive chart using Plotly."""
        if not PLOTLY_AVAILABLE:
            return "", False, "Plotly nao instalado. pip install plotly"
        
        try:
            labels = data.get("labels", data.get("x", []))
            values = data.get("values", data.get("y", []))
            
            if not labels or not values:
                return "", False, "Dados insuficientes (labels e values obrigatorios)"
            
            fig = None
            
            if chart_type == "bar":
                fig = px.bar(x=labels, y=values, title=title,
                             color_discrete_sequence=CORES_42C,
                             template="plotly_white")
                fig.update_layout(title_font_size=20, title_x=0.5)
                
            elif chart_type == "line":
                fig = px.line(x=labels, y=values, title=title,
                              markers=True, color_discrete_sequence=[COR_42C_SECUNDARIA],
                              template="plotly_white")
                
            elif chart_type == "pie":
                fig = px.pie(values=values, names=labels, title=title,
                             color_discrete_sequence=CORES_42C,
                             template="plotly_white")
                
            elif chart_type == "scatter":
                fig = px.scatter(x=labels, y=values, title=title,
                                 color_discrete_sequence=[COR_42C_PRIMARIA],
                                 template="plotly_white")
            
            elif chart_type == "area":
                fig = px.area(x=labels, y=values, title=title,
                              color_discrete_sequence=[COR_42C_ACAcento],
                              template="plotly_white")
            
            if fig is None:
                return "", False, f"Tipo de chart desconhecido: {chart_type}"
            
            # Salva como HTML interativo
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{chart_type}_{title.replace(' ', '_')}_{timestamp}.html"
            filepath = os.path.join(self.output_dir, filename)
            fig.write_html(filepath)
            
            # Tenta salvar PNG tambem (pode falhar sem Chrome)
            try:
                png_filename = f"{chart_type}_{title.replace(' ', '_')}_{timestamp}.png"
                png_path = os.path.join(self.output_dir, png_filename)
                fig.write_image(png_path, width=1200, height=700, scale=2)
            except Exception:
                pass  # HTML ja foi salvo, PNG e opcional
            
            return filepath, True, f"Plotly {chart_type} salvo: {filename}"
            
        except Exception as e:
            return "", False, f"Erro Plotly: {str(e)[:100]}"
    
    def generate_comparison_plotly(self, datasets: list, labels: list,
                                    title: str = "Comparacao") -> Tuple[str, bool, str]:
        """Multi-series comparison chart with Plotly."""
        if not PLOTLY_AVAILABLE:
            return "", False, "Plotly nao instalado"
        
        try:
            fig = go.Figure()
            for i, dataset in enumerate(datasets):
                name = dataset.get("name", f"Serie {i+1}")
                values = dataset.get("values", [])
                color = CORES_42C[i % len(CORES_42C)]
                fig.add_trace(go.Bar(name=name, x=labels, y=values, marker_color=color))
            
            fig.update_layout(
                title=title, title_font_size=20, title_x=0.5,
                barmode="group", template="plotly_white",
                legend_title_text=""
            )
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"comparison_{title.replace(' ', '_')}_{timestamp}.html"
            filepath = os.path.join(self.output_dir, filename)
            fig.write_html(filepath)
            
            return filepath, True, f"Plotly comparacao salvo: {filename}"
            
        except Exception as e:
            return "", False, f"Erro: {str(e)[:100]}"
    
    # =====================================================================
    # PPTX TEMPLATE 42c - Apresentacoes profissionais
    # =====================================================================
    
    
    
    def generate_pptx_42c(self, slides_content, title="Apresentacao 42c"):
        """Generate a professional PPTX with 42c branding."""
        if not PPTX_AVAILABLE:
            return "", False, "python-pptx nao instalado"
        
        try:
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            
            def _bg(slide, c):
                bg = slide.background
                fill = bg.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(int(c[1:3],16), int(c[3:5],16), int(c[5:7],16))
            
            def _tb(slide, l, t, w, h, text, fs=18, b=False, c="white", a=None):
                from pptx.util import Inches as I, Pt as P
                from pptx.enum.text import PP_ALIGN
                tx = slide.shapes.add_textbox(I(l), I(t), I(w), I(h))
                tf = tx.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = text
                p.font.size = P(fs)
                p.font.bold = b
                if c.startswith("#"):
                    p.font.color.rgb = RGBColor(int(c[1:3],16), int(c[3:5],16), int(c[5:7],16))
                else:
                    p.font.color.rgb = RGBColor(255,255,255)
                if a:
                    p.alignment = getattr(PP_ALIGN, a.upper(), None)
            
            from pptx.enum.shapes import MSO_SHAPE
            
            for sd in slides_content:
                st = sd.get("type", "title")
                stitle = sd.get("title", "")
                sc = sd.get("content", "")
                
                if st == "title":
                    sl = prs.slides.add_slide(prs.slide_layouts[6])
                    _bg(sl, "#1a1a2e")
                    _tb(sl, 1.5, 2.0, 10.3, 2, stitle, fs=44, b=True, a="CENTER")
                    if sc:
                        _tb(sl, 1.5, 4.0, 10.3, 1.5, sc, fs=22, a="CENTER")
                    ln = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(3.5), Inches(5.3), Inches(0.05))
                    ln.fill.solid()
                    ln.fill.fore_color.rgb = RGBColor(233, 69, 96)
                    ln.line.fill.background()
                
                elif st == "section":
                    sl = prs.slides.add_slide(prs.slide_layouts[6])
                    _bg(sl, "#0f3460")
                    _tb(sl, 1, 2.5, 11.3, 2, stitle, fs=40, b=True, a="CENTER")
                    if sc:
                        _tb(sl, 1, 4.2, 11.3, 1.5, sc, fs=20, a="CENTER")
                
                elif st == "bullet":
                    sl = prs.slides.add_slide(prs.slide_layouts[6])
                    _bg(sl, "#ffffff")
                    _tb(sl, 0.5, 0.3, 12.3, 0.8, stitle, fs=32, b=True, c="#1a1a2e")
                    ln = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.1), Inches(2), Inches(0.04))
                    ln.fill.solid()
                    ln.fill.fore_color.rgb = RGBColor(233, 69, 96)
                    ln.line.fill.background()
                    
                    if isinstance(sc, list):
                        tx = sl.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.5))
                        tf = tx.text_frame
                        tf.word_wrap = True
                        for i, item in enumerate(sc):
                            pp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                            pp.text = "  " + str(item)
                            pp.font.size = Pt(20)
                            pp.font.color.rgb = RGBColor(22, 33, 62)
                            pp.space_after = Pt(10)
                    elif sc:
                        _tb(sl, 0.8, 1.5, 11.5, 5, sc, fs=20, c="#16213e")
                
                elif st == "closing":
                    sl = prs.slides.add_slide(prs.slide_layouts[6])
                    _bg(sl, "#1a1a2e")
                    _tb(sl, 1.5, 2.5, 10.3, 1.5, stitle, fs=44, b=True, a="CENTER")
                    if sc:
                        _tb(sl, 1.5, 4.0, 10.3, 1, sc, fs=20, a="CENTER")
                    _tb(sl, 1.5, 5.5, 10.3, 0.5, "42c", fs=14, c="#e94560", a="CENTER")
            
            from datetime import datetime
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            fn = f"42c_{title.replace(chr(32), chr(95))}_{ts}.pptx"
            fp = os.path.join(self.output_dir, fn)
            prs.save(fp)
            return fp, True, f"PPTX 42c salvo: {fn}"
            
        except Exception as e:
            return "", False, f"Erro PPTX 42c: {str(e)[:100]}"
# Singleton
_viz_bridge = None

def get_viz_bridge() -> VizBridge:
    global _viz_bridge
    if _viz_bridge is None:
        _viz_bridge = VizBridge()
    return _viz_bridge
